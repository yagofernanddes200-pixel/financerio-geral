import streamlit as st
import pandas as pd
import plotly.express as px
import json
import os
from datetime import datetime

# Configuração da página para celular (Samsung A15 e iPhone)
st.set_page_config(
    page_title="Finanças Didáticas",
    page_icon="💰",
    layout="centered",
    initial_sidebar_state="collapsed"
)

# --- CONFIGURAÇÃO DE ESTILO: MODO ESCURO PREMIUM ---
st.markdown("""
<style>
    /* Estilo geral da página */
    .stApp {
        background-color: #121214;
        color: #E1E1E6;
    }
    
    /* Cabeçalho */
    h1, h2, h3, h4, h5, h6 {
        color: #FFFFFF !important;
        font-family: 'Inter', sans-serif;
    }
    
    /* Cards Customizados */
    .card-principal {
        background-color: #202024;
        padding: 18px;
        border-radius: 12px;
        margin-bottom: 15px;
        border-left: 5px solid #8257E5;
        box-shadow: 0 4px 6px rgba(0, 0, 0, 0.3);
    }
    .card-didatico {
        background-color: #1A1A1E;
        padding: 15px;
        border-radius: 10px;
        margin-bottom: 15px;
        border-left: 5px solid #3867D6;
        font-size: 0.9em;
        line-height: 1.5;
        color: #C4C4CC;
    }
    .card-receita {
        background-color: #1A2E26;
        padding: 15px;
        border-radius: 10px;
        border-left: 5px solid #26DE81;
    }
    .card-gasto {
        background-color: #2D1A1E;
        padding: 15px;
        border-radius: 10px;
        border-left: 5px solid #FF4757;
    }
    .titulo-secao {
        color: #FFFFFF;
        font-weight: bold;
        margin-top: 20px;
        border-bottom: 1px solid #29292E;
        padding-bottom: 8px;
    }
    
    /* Botões */
    .stButton>button {
        background-color: #8257E5 !important;
        color: white !important;
        border-radius: 8px !important;
        border: none !important;
        padding: 10px 20px !important;
        font-weight: bold !important;
        width: 100%;
        transition: background-color 0.2s;
    }
    .stButton>button:hover {
        background-color: #9466FF !important;
    }
    
    /* Inputs */
    div[data-baseweb="input"] {
        background-color: #202024 !important;
        border-color: #29292E !important;
    }
    input {
        color: #FFFFFF !important;
    }
</style>
""", unsafe_allow_html=True)

# --- FUNÇÕES DE AUXÍLIO PARA DATAS ---
def get_proximos_meses(data_inicial, parcelas):
    """Retorna uma lista de strings 'AAAA-MM' para as parcelas futuras."""
    datas = []
    ano = data_inicial.year
    mes = data_inicial.month
    
    for i in range(parcelas):
        datas.append(f"{ano}-{mes:02d}")
        mes += 1
        if mes > 12:
            mes = 1
            ano += 1
    return datas


# --- FUNÇÕES PARA GERAÇÃO DE RELATÓRIOS (PDF & WORD) ---
def gerar_relatorio_mensal_pdf(dados, periodo, mes_nome, ano, usuario_atual):
    from fpdf import FPDF
    from datetime import datetime
    
    # Filtrar dados do mês selecionado
    receitas_mes = [r for r in dados.get("receitas", []) if r["data"].startswith(periodo)]
    gastos_fixos_mes = [g for g in dados.get("gastos_fixos", []) if g["data"].startswith(periodo)]
    gastos_avulsos_mes = [g for g in dados.get("gastos_avulsos", []) if g["data"].startswith(periodo)]
    
    total_receitas = sum(r["valor"] for r in receitas_mes)
    total_fixos = sum(g["valor"] for g in gastos_fixos_mes)
    total_avulsos = sum(g["valor"] for g in gastos_avulsos_mes)
    total_gastos = total_fixos + total_avulsos
    saldo_livre = total_receitas - total_gastos
    
    # Economias do mês
    economias_mes = 0.0
    for t in dados.get("transferencias", []):
        if t["data"].startswith(periodo):
            dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
            orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
            if dest_tipo == "Guardado" and orig_tipo == "Normal":
                economias_mes += t["valor"]
                
    # Terceiros
    total_cartao_terceiros = sum(t["valor"] for t in dados.get("gastos_terceiros_cartao", []) if t["data"].startswith(periodo) and not t["pago"])
    total_emprestimos_pendentes = sum(e["valor"] for e in dados.get("gastos_terceiros_emprestimo", []) if not e["pago"])
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_margins(15, 15, 15)
    pdf.set_text_color(32, 32, 36)
    
    # Cabeçalho
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "FINANCAS DIDATICAS - RELATORIO MENSAL", ln=True, align="C")
    pdf.set_font("Helvetica", "", 12)
    pdf.cell(0, 8, f"Mes de Referencia: {mes_nome} de {ano}", ln=True, align="C")
    pdf.cell(0, 6, f"Usuario: {usuario_atual} | Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}", ln=True, align="C")
    pdf.ln(8)
    
    # Linha divisória Roxo
    pdf.set_draw_color(130, 87, 229)
    pdf.set_line_width(1)
    pdf.line(15, pdf.get_y(), 195, pdf.get_y())
    pdf.ln(6)
    
    # 1. Resumo Explicativo
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "1. RESUMO EXPLICATIVO DO PERIODO", ln=True)
    pdf.set_font("Helvetica", "", 10.5)
    
    texto = (
        f"Durante o mes de {mes_nome} de {ano}, o perfil '{usuario_atual}' apresentou as seguintes movimentacoes:\n\n"
        f"- Receitas Totais: Entrou um total de R$ {total_receitas:,.2f} em receitas cadastradas.\n"
        f"- Custos Totais: O total de despesas reais pagas ou a pagar no mes somou R$ {total_gastos:,.2f}, dividindo-se em R$ {total_fixos:,.2f} de gastos fixos/parcelamentos e R$ {total_avulsos:,.2f} de gastos avulsos (dia a dia).\n"
        f"- Balanco do Mes: O saldo livre apos abater as despesas foi de R$ {saldo_livre:,.2f}. "
    )
    if economias_mes > 0:
        texto += f"Desse saldo livre, R$ {economias_mes:,.2f} foram transferidos e salvos em contas poupanca/guardados (como o Banco Pan ou as Caixinhas do Nubank).\n"
    else:
        texto += "Nao foram registradas transferencias de economias para contas guardadas (poupanca) neste periodo.\n"
        
    texto += (
        f"\nEm relacao a terceiros (amigos/familiares):\n"
        f"- Cartao de Credito: Ha R$ {total_cartao_terceiros:,.2f} pendentes de reembolso em compras passadas no seu cartao para este mes.\n"
        f"- Emprestimos Diretos: O montante pendente de reembolso acumulado por emprestimos diretos (via PIX ou dinheiro) e de R$ {total_emprestimos_pendentes:,.2f}."
    )
    
    pdf.multi_cell(0, 5.5, texto.encode('latin-1', 'replace').decode('latin-1'))
    pdf.ln(8)
    
    # 2. Tabela de Indicadores
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "2. TABELA DE RESUMO MACRO DO MES", ln=True)
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    pdf.cell(100, 8, "Categoria de Fluxo", 1, 0, "L", fill=True)
    pdf.cell(80, 8, "Valor (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    indicadores = [
        ("Minhas Receitas (+)", total_receitas),
        ("Gastos Fixos (-)", total_fixos),
        ("Gastos Avulsos (-)", total_avulsos),
        ("Saldo Livre", saldo_livre),
        ("Economias Guardadas", economias_mes),
        ("Fatura de Terceiros a Receber (Mes)", total_cartao_terceiros),
        ("Emprestimos a Receber (Geral)", total_emprestimos_pendentes)
    ]
    for ind, val in indicadores:
        pdf.cell(100, 8, ind.encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(80, 8, f"R$ {val:,.2f}", 1, 1, "R")
        
    pdf.ln(8)
    
    # 3. Contas e Saldos
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "3. INTEGRACAO DE CONTAS E SALDOS", ln=True)
    pdf.set_font("Helvetica", "", 10.5)
    pdf.multi_cell(0, 5.5, "Abaixo estao listadas as contas cadastradas com seus saldos de fechamento atuais para conferencia:".encode('latin-1', 'replace').decode('latin-1'))
    pdf.ln(2)
    
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    pdf.cell(100, 8, "Conta / Banco", 1, 0, "L", fill=True)
    pdf.cell(40, 8, "Tipo", 1, 0, "C", fill=True)
    pdf.cell(40, 8, "Saldo Atual (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    for conta in dados.get("contas", []):
        pdf.cell(100, 8, conta["nome"].encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(40, 8, conta.get("tipo", "Normal"), 1, 0, "C")
        pdf.cell(40, 8, f"R$ {conta['saldo']:,.2f}", 1, 1, "R")
        
    return bytes(pdf.output())

def gerar_relatorio_mensal_docx(dados, periodo, mes_nome, ano, usuario_atual):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    import io
    
    # Filtrar dados do mês selecionado
    receitas_mes = [r for r in dados.get("receitas", []) if r["data"].startswith(periodo)]
    gastos_fixos_mes = [g for g in dados.get("gastos_fixos", []) if g["data"].startswith(periodo)]
    gastos_avulsos_mes = [g for g in dados.get("gastos_avulsos", []) if g["data"].startswith(periodo)]
    
    total_receitas = sum(r["valor"] for r in receitas_mes)
    total_fixos = sum(g["valor"] for g in gastos_fixos_mes)
    total_avulsos = sum(g["valor"] for g in gastos_avulsos_mes)
    total_gastos = total_fixos + total_avulsos
    saldo_livre = total_receitas - total_gastos
    
    # Economias
    economias_mes = 0.0
    for t in dados.get("transferencias", []):
        if t["data"].startswith(periodo):
            dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
            orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
            if dest_tipo == "Guardado" and orig_tipo == "Normal":
                economias_mes += t["valor"]
                
    # Terceiros
    total_cartao_terceiros = sum(t["valor"] for t in dados.get("gastos_terceiros_cartao", []) if t["data"].startswith(periodo) and not t["pago"])
    total_emprestimos_pendentes = sum(e["valor"] for e in dados.get("gastos_terceiros_emprestimo", []) if not e["pago"])
    
    doc = docx.Document()
    
    # Título do Relatório
    title = doc.add_paragraph()
    r_title = title.add_run(f"Relatório Financeiro Mensal - {mes_nome} / {ano}")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(130, 87, 229)
    
    subtitle = doc.add_paragraph()
    r_sub = subtitle.add_run(f"Finanças Didáticas | Usuário: {usuario_atual}\nGerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    
    doc.add_heading("1. Resumo Explicativo do Período", level=1)
    p1 = doc.add_paragraph()
    p1.add_run(f"Durante o mês de {mes_nome} de {ano}, o perfil '{usuario_atual}' realizou a gestão e organização do fluxo de caixa pessoal. Abaixo estão as análises detalhadas:\n\n")
    p1.add_run("Receitas Registradas: ").bold = True
    p1.add_run(f"Entrou um montante de R$ {total_receitas:,.2f}.\n")
    p1.add_run("Despesas do Mês: ").bold = True
    p1.add_run(f"As despesas totalizaram R$ {total_gastos:,.2f}, sendo compostas por R$ {total_fixos:,.2f} de custos fixos e parcelamentos, e R$ {total_avulsos:,.2f} de gastos avulsos (dia a dia).\n")
    p1.add_run("Saldo Livre: ").bold = True
    p1.add_run(f"O saldo livre restante (receitas menos despesas) foi de R$ {saldo_livre:,.2f}. ")
    if economias_mes > 0:
        p1.add_run(f"Deste saldo livre, você destinou R$ {economias_mes:,.2f} para suas contas de poupança (Guardados).\n")
    else:
        p1.add_run("Não foram registrados lançamentos de economias direcionadas para a poupança neste mês.\n")
        
    p1.add_run("\nSobre Finanças de Terceiros:\n").bold = True
    p1.add_run(f"- Cartão de Crédito: Há R$ {total_cartao_terceiros:,.2f} a serem reembolsados neste mês.\n")
    p1.add_run(f"- Empréstimos (PIX/Dinheiro): O total acumulado pendente de devolução a você é de R$ {total_emprestimos_pendentes:,.2f}.\n")
    
    doc.add_heading("2. Tabela de Indicadores Financeiros", level=1)
    
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.autofit = False
    
    # Header
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Indicador'
    hdr_cells[0].paragraphs[0].runs[0].font.bold = True
    hdr_cells[1].text = 'Valor (R$)'
    hdr_cells[1].paragraphs[0].runs[0].font.bold = True
    
    indicadores = [
        ("Minhas Receitas (+)", total_receitas),
        ("Gastos Fixos (-)", total_fixos),
        ("Gastos Avulsos (-)", total_avulsos),
        ("Saldo Livre", saldo_livre),
        ("Economias Guardadas", economias_mes),
        ("Fatura de Terceiros a Receber (Mes)", total_cartao_terceiros),
        ("Empréstimos a Receber (Geral)", total_emprestimos_pendentes)
    ]
    
    for ind, val in indicadores:
        row_cells = table.add_row().cells
        row_cells[0].text = ind
        row_cells[1].text = f"R$ {val:,.2f}"
        
    for row in table.rows:
        row.cells[0].width = Inches(4.5)
        row.cells[1].width = Inches(2.0)
        
    doc.add_paragraph("\n")
    doc.add_heading("3. Fechamento das Contas & Saldos", level=1)
    doc.add_paragraph("Abaixo estão listadas as contas e seus respectivos saldos de fechamento atuais:")
    
    table_contas = doc.add_table(rows=1, cols=3)
    table_contas.style = 'Table Grid'
    table_contas.autofit = False
    
    hdr_contas = table_contas.rows[0].cells
    hdr_contas[0].text = 'Conta / Banco'
    hdr_contas[0].paragraphs[0].runs[0].font.bold = True
    hdr_contas[1].text = 'Tipo'
    hdr_contas[1].paragraphs[0].runs[0].font.bold = True
    hdr_contas[2].text = 'Saldo Atual (R$)'
    hdr_contas[2].paragraphs[0].runs[0].font.bold = True
    
    for conta in dados.get("contas", []):
        row_cells = table_contas.add_row().cells
        row_cells[0].text = conta["nome"]
        row_cells[1].text = conta.get("tipo", "Normal")
        row_cells[2].text = f"R$ {conta['saldo']:,.2f}"
        
    for row in table_contas.rows:
        row.cells[0].width = Inches(3.0)
        row.cells[1].width = Inches(1.5)
        row.cells[2].width = Inches(2.0)
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def gerar_relatorio_anual_pdf(dados, ano, meses_lista, usuario_atual):
    from fpdf import FPDF
    from datetime import datetime
    
    dados_ano = []
    for m in range(1, 13):
        prefixo_busca = f"{ano}-{m:02d}"
        rec_ano = sum(r["valor"] for r in dados.get("receitas", []) if r["data"].startswith(prefixo_busca))
        fix_ano = sum(g["valor"] for g in dados.get("gastos_fixos", []) if g["data"].startswith(prefixo_busca))
        av_ano = sum(g["valor"] for g in dados.get("gastos_avulsos", []) if g["data"].startswith(prefixo_busca))
        gastos_totais = fix_ano + av_ano
        
        # Calcular economias
        economias_mes = 0.0
        for t in dados.get("transferencias", []):
            if t["data"].startswith(prefixo_busca):
                dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
                orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
                if dest_tipo == "Guardado" and orig_tipo == "Normal":
                    economias_mes += t["valor"]
                    
        dados_ano.append({
            "mes_num": m,
            "mes_nome": meses_lista[m-1],
            "ganhos": rec_ano,
            "gastos": gastos_totais,
            "guardado": economias_mes,
            "balanco": rec_ano - gastos_totais
        })
        
    total_ganhos_ano = sum(d["ganhos"] for d in dados_ano)
    total_gastos_ano = sum(d["gastos"] for d in dados_ano)
    total_guardado_ano = sum(d["guardado"] for d in dados_ano)
    saldo_anual_acumulado = total_ganhos_ano - total_gastos_ano
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_margins(15, 15, 15)
    pdf.set_text_color(32, 32, 36)
    
    # Cabeçalho
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "FINANCAS DIDATICAS - RELATORIO ANUAL CONSOLIDADO", ln=True, align="C")
    pdf.set_font("Helvetica", "", 12)
    pdf.cell(0, 8, f"Ano de Referencia: {ano}", ln=True, align="C")
    pdf.cell(0, 6, f"Usuario: {usuario_atual} | Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}", ln=True, align="C")
    pdf.ln(8)
    
    # Linha divisória Roxo
    pdf.set_draw_color(130, 87, 229)
    pdf.set_line_width(1)
    pdf.line(15, pdf.get_y(), 195, pdf.get_y())
    pdf.ln(6)
    
    # 1. Resumo Explicativo do Ano
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "1. RESUMO EXPLICATIVO DO ANO CONSOLIDADO", ln=True)
    pdf.set_font("Helvetica", "", 10.5)
    
    texto = (
        f"Este relatorio anual consolida a saude financeira do perfil '{usuario_atual}' ao longo de todo o ano de {ano}. "
        f"Abaixo estao as explicacoes e as metricas acumuladas do periodo:\n\n"
        f"- Ganhos Acumulados: Ao longo de todo o ano de {ano}, voce registrou um faturamento bruto acumulado de R$ {total_ganhos_ano:,.2f}.\n"
        f"- Despesas Acumuladas: O total de despesas reais somou R$ {total_gastos_ano:,.2f} (englobando todas as contas fixas, parcelamentos e gastos avulsos).\n"
        f"- Balanco Anual: O balanco financeiro liquido acumulado foi de R$ {saldo_anual_acumulado:,.2f}. "
    )
    if total_guardado_ano > 0:
        texto += f"Durante o ano, voce conseguiu guardar com sucesso um total de R$ {total_guardado_ano:,.2f} em economias direcionadas para contas poupanca/investimentos."
    else:
        texto += "Nao foram registradas transferencias de economias para poupancas de longo prazo durante o ano."
        
    pdf.multi_cell(0, 5.5, texto.encode('latin-1', 'replace').decode('latin-1'))
    pdf.ln(8)
    
    # 2. Tabela de Consolidação Mensal
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "2. TABELA DE COMPARAÇÃO MES A MES", ln=True)
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    
    pdf.cell(50, 8, "Mes", 1, 0, "L", fill=True)
    pdf.cell(40, 8, "Ganhos (R$)", 1, 0, "R", fill=True)
    pdf.cell(40, 8, "Gastos (R$)", 1, 0, "R", fill=True)
    pdf.cell(50, 8, "Balanco Liquido (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    for d in dados_ano:
        pdf.cell(50, 8, d["mes_nome"].encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(40, 8, f"R$ {d['ganhos']:,.2f}", 1, 0, "R")
        pdf.cell(40, 8, f"R$ {d['gastos']:,.2f}", 1, 0, "R")
        pdf.cell(50, 8, f"R$ {d['balanco']:,.2f}", 1, 1, "R")
        
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(230, 230, 235)
    pdf.cell(50, 8, "TOTAL ANUAL", 1, 0, "L", fill=True)
    pdf.cell(40, 8, f"R$ {total_ganhos_ano:,.2f}", 1, 0, "R", fill=True)
    pdf.cell(40, 8, f"R$ {total_gastos_ano:,.2f}", 1, 0, "R", fill=True)
    pdf.cell(50, 8, f"R$ {saldo_anual_acumulado:,.2f}", 1, 1, "R", fill=True)
    
    return bytes(pdf.output())

def gerar_relatorio_anual_docx(dados, ano, meses_lista, usuario_atual):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    import io
    from datetime import datetime
    
    dados_ano = []
    for m in range(1, 13):
        prefixo_busca = f"{ano}-{m:02d}"
        rec_ano = sum(r["valor"] for r in dados.get("receitas", []) if r["data"].startswith(prefixo_busca))
        fix_ano = sum(g["valor"] for g in dados.get("gastos_fixos", []) if g["data"].startswith(prefixo_busca))
        av_ano = sum(g["valor"] for g in dados.get("gastos_avulsos", []) if g["data"].startswith(prefixo_busca))
        gastos_totais = fix_ano + av_ano
        
        # Calcular economias
        economias_mes = 0.0
        for t in dados.get("transferencias", []):
            if t["data"].startswith(prefixo_busca):
                dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
                orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
                if dest_tipo == "Guardado" and orig_tipo == "Normal":
                    economias_mes += t["valor"]
                    
        dados_ano.append({
            "mes_nome": meses_lista[m-1],
            "ganhos": rec_ano,
            "gastos": gastos_totais,
            "guardado": economias_mes,
            "balanco": rec_ano - gastos_totais
        })
        
    total_ganhos_ano = sum(d["ganhos"] for d in dados_ano)
    total_gastos_ano = sum(d["gastos"] for d in dados_ano)
    total_guardado_ano = sum(d["guardado"] for d in dados_ano)
    saldo_anual_acumulado = total_ganhos_ano - total_gastos_ano
    
    doc = docx.Document()
    
    # Título do Relatório
    title = doc.add_paragraph()
    r_title = title.add_run(f"Relatório Financeiro Anual Consolidado - {ano}")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(130, 87, 229)
    
    subtitle = doc.add_paragraph()
    r_sub = subtitle.add_run(f"Finanças Didáticas | Usuário: {usuario_atual}\nGerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    
    doc.add_heading("1. Resumo Explicativo do Ano", level=1)
    p1 = doc.add_paragraph()
    p1.add_run(f"O presente relatório consolida o desempenho financeiro do perfil '{usuario_atual}' ao longo do ano de {ano}. A seguir, apresenta-se a síntese geral das movimentações:\n\n")
    p1.add_run("Ganhos Anuais Acumulados: ").bold = True
    p1.add_run(f"O faturamento bruto do ano somou R$ {total_ganhos_ano:,.2f}.\n")
    p1.add_run("Despesas Anuais Acumuladas: ").bold = True
    p1.add_run(f"O montante de saídas reais somou R$ {total_gastos_ano:,.2f}.\n")
    p1.add_run("Balanço Anual: ").bold = True
    p1.add_run(f"O balanço financeiro líquido acumulado foi de R$ {saldo_anual_acumulado:,.2f}. ")
    if total_guardado_ano > 0:
        p1.add_run(f"Deste valor, R$ {total_guardado_ano:,.2f} foram economizados com sucesso nas contas de poupança/investimentos (Guardados).\n")
    else:
        p1.add_run("Não foram registradas transferências de economias direcionadas para a poupança neste ano.\n")
        
    doc.add_heading("2. Detalhamento Mensal (Mês a Mês)", level=1)
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Mês'
    hdr_cells[0].paragraphs[0].runs[0].font.bold = True
    hdr_cells[1].text = 'Ganhos (R$)'
    hdr_cells[1].paragraphs[0].runs[0].font.bold = True
    hdr_cells[2].text = 'Gastos (R$)'
    hdr_cells[2].paragraphs[0].runs[0].font.bold = True
    hdr_cells[3].text = 'Balanço Líquido (R$)'
    hdr_cells[3].paragraphs[0].runs[0].font.bold = True
    
    for d in dados_ano:
        row_cells = table.add_row().cells
        row_cells[0].text = d["mes_nome"]
        row_cells[1].text = f"R$ {d['ganhos']:,.2f}"
        row_cells[2].text = f"R$ {d['gastos']:,.2f}"
        row_cells[3].text = f"R$ {d['balanco']:,.2f}"
        
    # Linha final de total
    total_row = table.add_row().cells
    total_row[0].text = "TOTAL ANUAL"
    total_row[0].paragraphs[0].runs[0].font.bold = True
    total_row[1].text = f"R$ {total_ganhos_ano:,.2f}"
    total_row[1].paragraphs[0].runs[0].font.bold = True
    total_row[2].text = f"R$ {total_gastos_ano:,.2f}"
    total_row[2].paragraphs[0].runs[0].font.bold = True
    total_row[3].text = f"R$ {saldo_anual_acumulado:,.2f}"
    total_row[3].paragraphs[0].runs[0].font.bold = True
    
    for row in table.rows:
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(1.5)
        row.cells[2].width = Inches(1.5)
        row.cells[3].width = Inches(1.5)
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
# --- CONTROLE DE MULTIPERFIS E SEGURANÇA ---
def obter_caminho_dados(usuario):
    nome_limpo = usuario.lower().strip().replace(" ", "_")
    return f"dados_financeiros_{nome_limpo}.json"

def carregar_dados_usuario(usuario):
    caminho = obter_caminho_dados(usuario)
    if os.path.exists(caminho):
        try:
            with open(caminho, "r", encoding="utf-8") as f:
                dados = json.load(f)
        except:
            dados = {}
    else:
        dados = {}
        
    # Inicialização padrão de chaves para segurança
    pin_padrao = "123456" if usuario.lower() == "yago" else "654321" if usuario.lower() == "binete" else "123456"
    dados.setdefault("pin", pin_padrao)
    dados.setdefault("receitas", [])
    dados.setdefault("gastos_fixos", [])
    dados.setdefault("gastos_avulsos", [])
    dados.setdefault("gastos_terceiros_cartao", [])
    dados.setdefault("gastos_terceiros_emprestimo", [])
    dados.setdefault("transferencias", [])
    
    # MIGRAR CONTAS ANTIGAS PARA ESTRUTURA DINÂMICA
    if "contas" not in dados:
        saldo_isolada = 0.0
        saldo_caixinha = 0.0
        if "conta_isolada" in dados:
            saldo_isolada = dados["conta_isolada"].get("saldo_principal", 0.0)
            saldo_caixinha = dados["conta_isolada"].get("caixinha_nu", 0.0)
            
        dados["contas"] = [
            {"nome": "Nu", "saldo": 0.0, "tipo": "Normal"},
            {"nome": "PicPay", "saldo": 0.0, "tipo": "Normal"},
            {"nome": "Inter", "saldo": 0.0, "tipo": "Normal"},
            {"nome": "Dinheiro", "saldo": 0.0, "tipo": "Normal"},
            {"nome": "Banco Pan (Guardado)", "saldo": saldo_isolada, "tipo": "Guardado"},
            {"nome": "Caixinha NU (Guardado)", "saldo": saldo_caixinha, "tipo": "Guardado"}
        ]
        
    return dados

def salvar_dados_usuario(usuario, dados):
    caminho = obter_caminho_dados(usuario)
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(dados, f, indent=4, ensure_ascii=False)

# --- AUXILIARES DE MOVIMENTAÇÃO DE SALDO ---
def alterar_saldo(dados, nome_conta, valor, operacao="somar"):
    for conta in dados["contas"]:
        if conta["nome"] == nome_conta:
            if operacao == "somar":
                conta["saldo"] += valor
            elif operacao == "subtrair":
                conta["saldo"] -= valor
            return True
    return False

# --- INICIALIZAÇÃO DO ESTADO DA SESSÃO ---
if "usuario_ativo" not in st.session_state:
    params = st.query_params
    usuario_url = params.get("user", "Yago").title()
    st.session_state.usuario_ativo = usuario_url

if "autenticado" not in st.session_state:
    st.session_state.autenticado = False

usuario_atual = st.session_state.usuario_ativo
dados = carregar_dados_usuario(usuario_atual)

# --- CABEÇALHO DO APP ---
col_logo, col_titulo = st.columns([1, 4])
with col_logo:
    if os.path.exists("app_logo_financeiro.png"):
        st.image("app_logo_financeiro.png", width=70)
    else:
        st.markdown("<h1 style='font-size: 50px; margin: 0;'>💰</h1>", unsafe_allow_html=True)

with col_titulo:
    st.title("Finanças Didáticas")
    st.caption("Controle financeiro premium e descomplicado")

# --- TELA DE LOGIN / SEGURANÇA ---
if not st.session_state.autenticado:
    st.markdown(f"### 🔒 Acesso Restrito - Perfil: **{usuario_atual}**")
    st.markdown("Por segurança, digite sua senha de 6 dígitos para acessar seus dados.")
    
    pin_digitado = st.text_input("Digite o PIN de 6 dígitos:", type="password", max_chars=6)
    
    col_btn_login, col_btn_perfil = st.columns(2)
    with col_btn_login:
        if st.button("Entrar 🔓"):
            if pin_digitado == dados.get("pin", "123456"):
                st.session_state.autenticado = True
                st.success("Acesso liberado!")
                st.rerun()
            else:
                st.error("PIN incorreto! Tente novamente.")
                
    with col_btn_perfil:
        novo_usuario = st.selectbox("Trocar de Perfil:", ["Yago", "Binete", "Criar Novo Perfil"])
        if novo_usuario == "Criar Novo Perfil":
            st.markdown("---")
            st.markdown("**Criar Novo Perfil Independente**")
            novo_nome = st.text_input("Nome do Novo Usuário:")
            novo_pin = st.text_input("Escolha um PIN de 6 dígitos (somente números):", type="password", max_chars=6)
            if st.button("Criar Perfil e Acessar 🚀"):
                if novo_nome and len(novo_pin) == 6 and novo_pin.isdigit():
                    dados_novos = carregar_dados_usuario(novo_nome)
                    dados_novos["pin"] = novo_pin
                    salvar_dados_usuario(novo_nome, dados_novos)
                    st.session_state.usuario_ativo = novo_nome.title()
                    st.session_state.autenticado = True
                    st.query_params["user"] = novo_nome.lower()
                    st.success(f"Perfil de {novo_nome} criado e logado!")
                    st.rerun()
                else:
                    st.error("Por favor, digite um nome válido e um PIN numérico de exatamente 6 dígitos.")
        elif novo_usuario != usuario_atual:
            st.session_state.usuario_ativo = novo_usuario
            st.query_params["user"] = novo_usuario.lower()
            st.rerun()
            
    st.stop()

# --- SELETOR DE MÊS E ANO (PLANEJAMENTO SEM LIMITES) ---
st.markdown("---")
col_perf, col_m, col_a = st.columns([2, 2, 2])

with col_perf:
    st.markdown(f"👤 **{usuario_atual}**")
    if st.button("Sair / Bloquear 🔒"):
        st.session_state.autenticado = False
        st.rerun()

with col_m:
    meses_lista = [
        "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
        "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"
    ]
    mes_atual_nome = meses_lista[datetime.now().month - 1]
    mes_selecionado = st.selectbox("Selecione o Mês:", meses_lista, index=datetime.now().month - 1)
    mes_num = meses_lista.index(mes_selecionado) + 1

with col_a:
    anos_lista = list(range(2025, 2036))
    ano_selecionado = st.selectbox("Selecione o Ano:", anos_lista, index=anos_lista.index(datetime.now().year))

periodo_ativo = f"{ano_selecionado}-{mes_num:02d}"

# --- NAVEGAÇÃO POR ABAS (Mobile Friendly) ---
abas = st.tabs(["📊 Resumo", "💵 Ganhos", "🏠 Fixas", "🛍️ Avulsos", "👥 Terceiros", "💳 Saldos", "🔒 Guardado", "⚙️ Senha"])

# ----------------- ABA 1: RESUMO GERAL & RELATÓRIOS -----------------
with abas[0]:
    st.markdown(f"<h3 class='titulo-secao'>📊 Painel Financeiro ({mes_selecionado}/{ano_selecionado})</h3>", unsafe_allow_html=True)
    
    tipo_relatorio = st.radio("Escolha o tipo de Relatório:", ["Relatório Mensal", "Relatório Anual"], horizontal=True)
    
    if tipo_relatorio == "Relatório Mensal":
        # Filtrar dados do mês selecionado
        receitas_mes = [r for r in dados.get("receitas", []) if r["data"].startswith(periodo_ativo)]
        gastos_fixos_mes = [g for g in dados.get("gastos_fixos", []) if g["data"].startswith(periodo_ativo)]
        gastos_avulsos_mes = [g for g in dados.get("gastos_avulsos", []) if g["data"].startswith(periodo_ativo)]
        
        total_receitas = sum(r["valor"] for r in receitas_mes)
        total_fixos = sum(g["valor"] for g in gastos_fixos_mes)
        total_avulsos = sum(g["valor"] for g in gastos_avulsos_mes)
        total_gastos = total_fixos + total_avulsos
        saldo_livre = total_receitas - total_gastos
        
        # Calcular total guardado nas contas do tipo "Guardado"
        total_guardado = sum(c["saldo"] for c in dados["contas"] if c["tipo"] == "Guardado")
        
        # Terceiros pendentes no cartão do mês
        total_cartao_terceiros = sum(t["valor"] for t in dados.get("gastos_terceiros_cartao", []) if t["data"].startswith(periodo_ativo) and not t["pago"])
        # Empréstimos acumulados (Dinheiro/PIX) - pendentes globais
        total_emprestimos_pendentes = sum(e["valor"] for e in dados.get("gastos_terceiros_emprestimo", []) if not e["pago"])
        
        # Métricas na tela
        col_m1, col_m2 = st.columns(2)
        with col_m1:
            st.metric(label="Minhas Receitas (+)", value=f"R$ {total_receitas:,.2f}")
            st.metric(label="Meus Gastos (-)", value=f"R$ {total_gastos:,.2f}")
        with col_m2:
            st.metric(label="Saldo do Mês (Livre)", value=f"R$ {saldo_livre:,.2f}", delta=f"R$ {saldo_livre:,.2f}" if saldo_livre >= 0 else f"R$ {saldo_livre:,.2f}")
            st.metric(label="Total Guardado", value=f"R$ {total_guardado:,.2f}")
            
        # Alertas de Terceiros
        st.markdown("#### 👥 Dinheiro de Terceiros a Receber:")
        c_t1, c_t2 = st.columns(2)
        with c_t1:
            if total_cartao_terceiros > 0:
                st.warning(f"💳 **Cartão de Crédito:**\n\nR$ {total_cartao_terceiros:,.2f} pendentes neste mês.")
            else:
                st.success("💳 **Cartão de Crédito:**\n\nNenhum reembolso pendente.")
        with c_t2:
            if total_emprestimos_pendentes > 0:
                st.warning(f"💸 **Empréstimos:**\n\nR$ {total_emprestimos_pendentes:,.2f} a receber.")
            else:
                st.success("💸 **Empréstimos:**\n\nTodos pagos!")
                
        # Gráficos do Mês
        col_graf_1, col_graf_2 = st.columns(2) if not st.session_state.get("mobile", True) else (st.container(), st.container())
        
        # Gráfico 1: Divisão das Despesas (Pizza)
        with col_graf_1:
            if total_gastos > 0:
                st.markdown("#### 📈 Divisão das Despesas (Fixos vs. Avulsos)")
                df_pizza = pd.DataFrame([
                    {"Categoria": "Gastos Fixos", "Valor": total_fixos},
                    {"Categoria": "Gastos Avulsos", "Valor": total_avulsos}
                ])
                fig = px.pie(df_pizza, values="Valor", names="Categoria", hole=0.4,
                             color_discrete_sequence=["#8257E5", "#FF4757"])
                fig.update_layout(
                    paper_bgcolor='rgba(0,0,0,0)',
                    plot_bgcolor='rgba(0,0,0,0)',
                    font_color="#FFFFFF",
                    margin=dict(t=10, b=10, l=10, r=10),
                    height=220
                )
                st.plotly_chart(fig, use_container_width=True)
            
        # Gráfico 2: Alocação Geral Macro (Ganhos vs. Despesas vs. Guardados do Mês)
        with col_graf_2:
            # Calcula depósitos para poupança neste mês
            economias_mes = 0.0
            for t in dados.get("transferencias", []):
                if t["data"].startswith(periodo_ativo):
                    dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
                    orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
                    if dest_tipo == "Guardado" and orig_tipo == "Normal":
                        economias_mes += t["valor"]
                        
            st.markdown("#### 📊 Alocação Macro Mensal (Fluxos)")
            df_macro = pd.DataFrame([
                {"Fluxo": "Ganhos", "Valor": total_receitas, "Cor": "#26DE81"},
                {"Fluxo": "Despesas", "Valor": total_gastos, "Cor": "#FF4757"},
                {"Fluxo": "Guardado", "Valor": economias_mes, "Cor": "#8257E5"}
            ])
            fig_macro = px.bar(df_macro, x="Valor", y="Fluxo", color="Fluxo", orientation="h",
                               color_discrete_map={"Ganhos": "#26DE81", "Despesas": "#FF4757", "Guardado": "#8257E5"})
            fig_macro.update_layout(
                paper_bgcolor='rgba(0,0,0,0)',
                plot_bgcolor='rgba(0,0,0,0)',
                font_color="#FFFFFF",
                margin=dict(t=10, b=10, l=10, r=10),
                height=200,
                showlegend=False
            )
            st.plotly_chart(fig_macro, use_container_width=True)
            

        # --- EXPORTAR RELATÓRIO MENSAL (BOTÕES) ---
        st.markdown("---")
        st.markdown("### 📥 Exportar Relatório Oficial (Mensal)")
        st.caption("Gere um relatório explicativo completo, pronto para baixar em PDF ou editar em Word.")
        
        col_pdf, col_docx = st.columns(2)
        with col_pdf:
            pdf_bytes = gerar_relatorio_mensal_pdf(dados, periodo_ativo, mes_selecionado, ano_selecionado, usuario_atual)
            st.download_button(
                label="📄 Baixar PDF",
                data=pdf_bytes,
                file_name=f"relatorio_financeiro_{usuario_atual.lower()}_{periodo_ativo}.pdf",
                mime="application/pdf"
            )
        with col_docx:
            docx_bytes = gerar_relatorio_mensal_docx(dados, periodo_ativo, mes_selecionado, ano_selecionado, usuario_atual)
            st.download_button(
                label="📝 Baixar Word (.docx)",
                data=docx_bytes,
                file_name=f"relatorio_financeiro_{usuario_atual.lower()}_{periodo_ativo}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
    else:
        # --- RELATÓRIO ANUAL CONSOLIDADO ---
        st.markdown(f"### 📅 Relatório Anual - Ano: **{ano_selecionado}**")
        
        # Consolidar dados de todos os 12 meses do ano selecionado de forma correta
        dados_ano = []
        for m in range(1, 13):
            prefixo_busca = f"{ano_selecionado}-{m:02d}"
            rec_ano = sum(r["valor"] for r in dados.get("receitas", []) if r["data"].startswith(prefixo_busca))
            fix_ano = sum(g["valor"] for g in dados.get("gastos_fixos", []) if g["data"].startswith(prefixo_busca))
            av_ano = sum(g["valor"] for g in dados.get("gastos_avulsos", []) if g["data"].startswith(prefixo_busca)) # CORRIGIDO O BUG AQUI!
            gastos_totais = fix_ano + av_ano
            dados_ano.append({
                "Mês": meses_lista[m-1][:3],
                "Ganhos": rec_ano,
                "Gastos": gastos_totais
            })
            
        df_ano = pd.DataFrame(dados_ano)
        
        # Totais Anuais
        total_ganhos_ano = df_ano["Ganhos"].sum()
        total_gastos_ano = df_ano["Gastos"].sum()
        saldo_anual_acumulado = total_ganhos_ano - total_gastos_ano
        
        col_an1, col_an2, col_an3 = st.columns(3)
        col_an1.metric("Ganhos no Ano", f"R$ {total_ganhos_ano:,.2f}")
        col_an2.metric("Gastos no Ano", f"R$ {total_gastos_ano:,.2f}")
        col_an3.metric("Balanço Acumulado", f"R$ {saldo_anual_acumulado:,.2f}", 
                        delta=f"R$ {saldo_anual_acumulado:,.2f}" if saldo_anual_acumulado >= 0 else f"R$ {saldo_anual_acumulado:,.2f}")
        
        # Gráfico de Linha/Tendência Premium (Substituindo as barras como Yago pediu)
        st.markdown("#### 📈 Tendência Anual (Ganhos vs. Gastos)")
        df_melted = df_ano.melt(id_vars="Mês", value_vars=["Ganhos", "Gastos"], var_name="Categoria", value_name="Valor (R$)")
        fig_trend = px.line(df_melted, x="Mês", y="Valor (R$)", color="Categoria", markers=True,
                            color_discrete_map={"Ganhos": "#26DE81", "Gastos": "#FF4757"})
        fig_trend.update_layout(
            paper_bgcolor='rgba(0,0,0,0)',
            plot_bgcolor='rgba(0,0,0,0)',
            font_color="#FFFFFF",
            margin=dict(t=10, b=10, l=10, r=10),
            height=280
        )
        st.plotly_chart(fig_trend, use_container_width=True)


        # --- EXPORTAR RELATÓRIO ANUAL (BOTÕES) ---
        st.markdown("---")
        st.markdown("### 📥 Exportar Relatório Oficial (Anual)")
        st.caption("Gere um relatório consolidado de todo o ano, com análises explicativas, tabelas e comparativos.")
        
        col_pdf, col_docx = st.columns(2)
        with col_pdf:
            pdf_bytes_anual = gerar_relatorio_anual_pdf(dados, ano_selecionado, meses_lista, usuario_atual)
            st.download_button(
                label="📄 Baixar PDF Anual",
                data=pdf_bytes_anual,
                file_name=f"relatorio_anual_{usuario_atual.lower()}_{ano_selecionado}.pdf",
                mime="application/pdf"
            )
        with col_docx:
            docx_bytes_anual = gerar_relatorio_anual_docx(dados, ano_selecionado, meses_lista, usuario_atual)
            st.download_button(
                label="📝 Baixar Word Anual (.docx)",
                data=docx_bytes_anual,
                file_name=f"relatorio_anual_{usuario_atual.lower()}_{ano_selecionado}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
# ----------------- ABA 2: GANHOS (RECEITAS) -----------------
with abas[1]:
    st.markdown("<h3 class='titulo-secao'>💵 Registro de Ganhos</h3>", unsafe_allow_html=True)
    
    with st.form("form_receita", clear_on_submit=True):
        st.markdown("**Adicionar Nova Receita**")
        desc = st.text_input("Descrição (Ex: Salário, PIX, Freelance)")
        val = st.number_input("Valor (R$)", min_value=0.0, step=50.0, format="%.2f")
        
        # Selecionar em qual conta depositar
        contas_nomes = [c["nome"] for c in dados["contas"]]
        conta_dest = st.selectbox("Depositar na Conta:", contas_nomes)
        
        data_padrao = datetime(ano_selecionado, mes_num, min(datetime.now().day, 28))
        data_rec = st.date_input("Data do Recebimento", data_padrao)
        
        enviar = st.form_submit_button("Salvar Receita")
        if enviar and desc and val > 0:
            nova_rec = {"data": str(data_rec), "descricao": desc, "valor": val, "conta": conta_dest}
            dados.setdefault("receitas", []).append(nova_rec)
            
            # Atualiza saldo automaticamente
            alterar_saldo(dados, conta_dest, val, "somar")
            
            salvar_dados_usuario(usuario_atual, dados)
            st.success(f"Receita adicionada e valor depositado na conta {conta_dest}!")
            st.rerun()

    # Mostrar Receitas do Mês Selecionado
    receitas_mes = [r for r in dados.get("receitas", []) if r["data"].startswith(periodo_ativo)]
    if receitas_mes:
        st.markdown("#### Suas Receitas Registradas neste Mês")
        df_rec = pd.DataFrame(receitas_mes)
        # Se a coluna 'conta' não existir nos dados antigos, preencher com 'Não especificada'
        if "conta" not in df_rec.columns:
            df_rec["conta"] = "Não especificada"
        df_rec = df_rec[["data", "descricao", "conta", "valor"]]
        df_rec.columns = ["Data", "Descrição", "Conta", "Valor (R$)"]
        st.dataframe(df_rec, use_container_width=True)
        
        if st.button("Limpar Histórico de Receitas (Mês)", key="limpar_rec"):
            # Devolver o valor ao limpar? Para simplificar, limpa apenas o histórico visual de lançamentos do mês
            dados["receitas"] = [r for r in dados.get("receitas", []) if not r["data"].startswith(periodo_ativo)]
            salvar_dados_usuario(usuario_atual, dados)
            st.rerun()

# ----------------- ABA 3: MEUS GASTOS FIXOS -----------------
with abas[2]:
    st.markdown("<h3 class='titulo-secao'>🏠 Meus Gastos Fixos & Parcelas</h3>", unsafe_allow_html=True)
    st.caption("Gastos repetitivos ou parcelados de longo prazo.")

    with st.form("form_fixo", clear_on_submit=True):
        st.markdown("**Adicionar Gasto Fixo / Parcelado**")
        desc = st.text_input("Descrição do Gasto (Ex: Aluguel, Parcela de Notebook)")
        val = st.number_input("Valor Mensal (R$)", min_value=0.0, step=10.0, format="%.2f")
        
        contas_nomes = [c["nome"] for c in dados["contas"]]
        conta_pagamento = st.selectbox("Pagar com a Conta:", contas_nomes)
        
        is_parcelado = st.checkbox("Esta despesa é parcelada?")
        total_parc = st.number_input("Número total de parcelas:", min_value=1, max_value=48, value=1, step=1)
        pago = st.checkbox("Marcar primeira parcela como já paga?")
        
        enviar = st.form_submit_button("Salvar Despesa")
        if enviar and desc and val > 0:
            data_inicial = datetime(ano_selecionado, mes_num, 1)
            
            if is_parcelado:
                lista_periodos = get_proximos_meses(data_inicial, total_parc)
                for i, periodo in enumerate(lista_periodos):
                    status_pago_parc = pago if i == 0 else False
                    novo_gasto = {
                        "data": f"{periodo}-01",
                        "descricao": f"{desc} (Parc. {i+1}/{total_parc})",
                        "valor": val,
                        "pago": status_pago_parc,
                        "conta": conta_pagamento
                    }
                    dados.setdefault("gastos_fixos", []).append(novo_gasto)
                    if status_pago_parc:
                        alterar_saldo(dados, conta_pagamento, val, "subtrair")
            else:
                novo_gasto = {
                    "data": f"{periodo_ativo}-01",
                    "descricao": desc,
                    "valor": val,
                    "pago": pago,
                    "conta": conta_pagamento
                }
                dados.setdefault("gastos_fixos", []).append(novo_gasto)
                if pago:
                    alterar_saldo(dados, conta_pagamento, val, "subtrair")
                
            salvar_dados_usuario(usuario_atual, dados)
            st.success("Despesa cadastrada com sucesso!")
            st.rerun()

    # Mostrar Gastos Fixos do Mês
    gastos_fixos_mes = [g for g in dados.get("gastos_fixos", []) if g["data"].startswith(periodo_ativo)]
    if gastos_fixos_mes:
        st.markdown("#### Controle de Pagamentos de Despesas Fixas")
        for idx, item in enumerate(dados["gastos_fixos"]):
            if item["data"].startswith(periodo_ativo):
                col_d, col_v, col_p = st.columns([2, 1, 1])
                conta_item = item.get("conta", "Nu")
                col_d.markdown(f"**{item['descricao']}**\n\n*({conta_item})*")
                col_v.markdown(f"R$ {item['valor']:,.2f}")
                
                status_pago = col_p.checkbox("Pago", value=item["pago"], key=f"fixo_{idx}_{periodo_ativo}")
                if status_pago != item["pago"]:
                    # Se mudou o status, atualizar saldo da conta correspondente
                    item_conta = item.get("conta", "Nu")
                    if status_pago: # Foi pago agora
                        alterar_saldo(dados, item_conta, item["valor"], "subtrair")
                    else: # Foi desmarcado como pago, devolve saldo
                        alterar_saldo(dados, item_conta, item["valor"], "somar")
                        
                    dados["gastos_fixos"][idx]["pago"] = status_pago
                    salvar_dados_usuario(usuario_atual, dados)
                    st.rerun()
                    
        if st.button("Limpar Todos os Gastos Fixos (Deste Mês)", key="limpar_fixos"):
            dados["gastos_fixos"] = [g for g in dados.get("gastos_fixos", []) if not g["data"].startswith(periodo_ativo)]
            salvar_dados_usuario(usuario_atual, dados)
            st.rerun()

# ----------------- ABA 4: GASTOS AVULSOS -----------------
with abas[3]:
    st.markdown("<h3 class='titulo-secao'>🛍️ Meus Gastos Avulsos</h3>", unsafe_allow_html=True)
    st.caption("Gastos diários variáveis (Ifood, Uber, Compras rápidas)")

    with st.form("form_avulso", clear_on_submit=True):
        st.markdown("**Adicionar Gasto Avulso**")
        desc = st.text_input("O que você comprou?")
        val = st.number_input("Valor Pago (R$)", min_value=0.0, step=10.0, format="%.2f")
        
        contas_nomes = [c["nome"] for c in dados["contas"]]
        conta_pagamento = st.selectbox("Pagar com a Conta:", contas_nomes)
        
        data_padrao = datetime(ano_selecionado, mes_num, min(datetime.now().day, 28))
        data_gasto = st.date_input("Data do Gasto", data_padrao)
        
        enviar = st.form_submit_button("Salvar Gasto Avulso")
        if enviar and desc and val > 0:
            novo_avulso = {"data": str(data_gasto), "descricao": desc, "valor": val, "conta": conta_pagamento}
            dados.setdefault("gastos_avulsos", []).append(novo_avulso)
            
            # Deduz do saldo automaticamente
            alterar_saldo(dados, conta_pagamento, val, "subtrair")
            
            salvar_dados_usuario(usuario_atual, dados)
            st.success(f"Gasto avulso adicionado e deduzido da conta {conta_pagamento}!")
            st.rerun()

    # Mostrar Gastos Avulsos
    gastos_avulsos_mes = [g for g in dados.get("gastos_avulsos", []) if g["data"].startswith(periodo_ativo)]
    if gastos_avulsos_mes:
        st.markdown("#### Seus Gastos Avulsos Registrados")
        df_av = pd.DataFrame(gastos_avulsos_mes)
        if "conta" not in df_av.columns:
            df_av["conta"] = "Não especificada"
        df_av = df_av[["data", "descricao", "conta", "valor"]]
        df_av.columns = ["Data", "Descrição", "Conta", "Valor (R$)"]
        st.dataframe(df_av, use_container_width=True)
        
        if st.button("Limpar Histórico de Avulsos (Mês)", key="limpar_avulsos"):
            dados["gastos_avulsos"] = [g for g in dados.get("gastos_avulsos", []) if not g["data"].startswith(periodo_ativo)]
            salvar_dados_usuario(usuario_atual, dados)
            st.rerun()

# ----------------- ABA 5: GASTOS DE TERCEIROS (CARTÃO VS EMPRÉSTIMO) -----------------
with abas[4]:
    st.markdown("<h3 class='titulo-secao'>👥 Amigos & Terceiros no Meu Cartão / Empréstimos</h3>", unsafe_allow_html=True)
    
    sub_abas_terceiros = st.tabs(["💳 Compras no Meu Cartão", "💸 Empréstimos (PIX/Dinheiro)"])
    
    # SUB-ABA 5.1: COMPRAS NO MEU CARTÃO
    with sub_abas_terceiros[0]:
        st.markdown("**Cadastro de Compras de Terceiros no seu Cartão de Crédito**")
        st.caption("Quando alguém usa seu cartão para comprar parcelado ou à vista.")
        
        with st.form("form_cartao_terceiros", clear_on_submit=True):
            nome = st.text_input("Nome de quem comprou:")
            desc = st.text_input("Descrição da compra (Ex: Pizza, Tênis)")
            val = st.number_input("Valor da compra/parcela (R$)", min_value=0.0, step=10.0, format="%.2f")
            
            is_parc_terc = st.checkbox("Esta compra de terceiro é parcelada?")
            total_parc_terc = st.number_input("Número de parcelas de terceiro:", min_value=1, max_value=48, value=1, step=1)
            pago_terc = st.checkbox("Já te pagou a primeira parcela?")
            
            enviar_c = st.form_submit_button("Salvar Compra no Cartão")
            if enviar_c and nome and desc and val > 0:
                data_inicial = datetime(ano_selecionado, mes_num, 1)
                
                if is_parc_terc:
                    lista_periodos = get_proximos_meses(data_inicial, total_parc_terc)
                    for i, periodo in enumerate(lista_periodos):
                        nova_compra = {
                            "data": f"{periodo}-01",
                            "nome": nome.title(),
                            "descricao": f"{desc} (Parc. {i+1}/{total_parc_terc})",
                            "valor": val,
                            "pago": pago_terc if i == 0 else False
                        }
                        dados.setdefault("gastos_terceiros_cartao", []).append(nova_compra)
                else:
                    nova_compra = {
                        "data": f"{periodo_ativo}-01",
                        "nome": nome.title(),
                        "descricao": desc,
                        "valor": val,
                        "pago": pago_terc
                    }
                    dados.setdefault("gastos_terceiros_cartao", []).append(nova_compra)
                    
                salvar_dados_usuario(usuario_atual, dados)
                st.success("Compra no cartão adicionada com sucesso!")
                st.rerun()

        # Visualizar faturas do mês de terceiros
        cartao_mes = [t for t in dados.get("gastos_terceiros_cartao", []) if t["data"].startswith(periodo_ativo)]
        if cartao_mes:
            st.markdown(f"#### Reembolsos de Fatura de Terceiros - {mes_selecionado}/{ano_selecionado}")
            for idx, item in enumerate(dados["gastos_terceiros_cartao"]):
                if item["data"].startswith(periodo_ativo):
                    c_det, c_status = st.columns([3, 1])
                    status_txt = "✅ Recebido" if item["pago"] else "❌ Pendente"
                    c_det.markdown(f"**{item['nome']}** - {item['descricao']} (R$ {item['valor']:,.2f}) [{status_txt}]")
                    
                    pago_atualizado = c_status.checkbox("Recebido", value=item["pago"], key=f"cartao_terc_{idx}_{periodo_ativo}")
                    if pago_atualizado != item["pago"]:
                        dados["gastos_terceiros_cartao"][idx]["pago"] = pago_atualizado
                        salvar_dados_usuario(usuario_atual, dados)
                        st.rerun()
                        
    # SUB-ABA 5.2: EMPRÉSTIMOS (PIX/DINHEIRO) - INTEGRADO COM CONTAS
    with sub_abas_terceiros[1]:
        st.markdown("**Cadastro de Empréstimos (Dinheiro físico ou PIX)**")
        st.caption("Dinheiro que saiu de uma conta sua e que deve retornar para alguma conta sua.")
        
        with st.form("form_emprestimo_terceiros", clear_on_submit=True):
            nome_emp = st.text_input("Quem pediu o empréstimo?")
            desc_emp = st.text_input("Motivo do empréstimo (Ex: Empréstimo para Luz)")
            val_emp = st.number_input("Valor Emprestado (R$)", min_value=0.0, step=10.0, format="%.2f")
            
            # Selecionar de qual conta tirar o dinheiro
            contas_nomes = [c["nome"] for c in dados["contas"]]
            conta_origem_emp = st.selectbox("Retirar Dinheiro de qual Conta/Dinheiro?", contas_nomes)
            
            enviar_e = st.form_submit_button("Registrar Empréstimo")
            if enviar_e and nome_emp and val_emp > 0:
                novo_emp = {
                    "data": str(datetime.now().date()),
                    "nome": nome_emp.title(),
                    "descricao": desc_emp,
                    "valor": val_emp,
                    "pago": False,
                    "conta_origem": conta_origem_emp,
                    "forma_recebimento": "",
                    "conta_destino": ""
                }
                dados.setdefault("gastos_terceiros_emprestimo", []).append(novo_emp)
                
                # Deduz da conta de origem automaticamente
                alterar_saldo(dados, conta_origem_emp, val_emp, "subtrair")
                
                salvar_dados_usuario(usuario_atual, dados)
                st.success(f"Empréstimo registrado! R$ {val_emp:,.2f} foi deduzido de '{conta_origem_emp}'.")
                st.rerun()

        # Lista total de empréstimos pendentes
        emprestimos_todos = dados.get("gastos_terceiros_emprestimo", [])
        if emprestimos_todos:
            st.markdown("#### Histórico Geral de Empréstimos")
            
            # Resumo por pessoa
            divida_pessoa = {}
            for e in emprestimos_todos:
                if not e["pago"]:
                    divida_pessoa[e["nome"]] = divida_pessoa.get(e["nome"], 0.0) + e["valor"]
                    
            if divida_pessoa:
                st.markdown("**Saldos de Devedores de Empréstimo:**")
                for pessoa, valor in divida_pessoa.items():
                    st.info(f"👤 **{pessoa}** te deve um total de: **R$ {valor:,.2f}**")
                    
            st.markdown("---")
            st.markdown("**Atualizar Reembolso / Recebimento de Empréstimo:**")
            for idx, item in enumerate(dados["gastos_terceiros_emprestimo"]):
                col_e_det, col_e_st = st.columns([3, 1])
                status_txt = "✅ Pago" if item["pago"] else "❌ Pendente"
                origem_item = item.get("conta_origem", "Não especificada")
                
                # Detalhes textuais ricos do empréstimo
                detalhe_texto = f"**{item['nome']}** - {item['descricao']} (R$ {item['valor']:,.2f})\n\n*(Origem: {origem_item})*"
                if item["pago"]:
                    detalhe_texto += f" \n\n*(Recebido em {item.get('forma_recebimento', 'PIX')} na conta {item.get('conta_destino', 'Nu')})*"
                col_e_det.markdown(detalhe_texto)
                
                # Interatividade para recebimento
                if not item["pago"]:
                    with col_e_st.popover("Receber"):
                        st.markdown("**Confirmar Recebimento**")
                        forma_rec = st.selectbox("Forma:", ["PIX", "Dinheiro"], key=f"forma_{idx}")
                        contas_nomes = [c["nome"] for c in dados["contas"]]
                        conta_dest_reemb = st.selectbox("Depositar na Conta:", contas_nomes, key=f"dest_{idx}")
                        
                        if st.button("Confirmar", key=f"conf_rec_{idx}"):
                            dados["gastos_terceiros_emprestimo"][idx]["pago"] = True
                            dados["gastos_terceiros_emprestimo"][idx]["forma_recebimento"] = forma_rec
                            dados["gastos_terceiros_emprestimo"][idx]["conta_destino"] = conta_dest_reemb
                            
                            # Adiciona de volta ao saldo da conta escolhida
                            alterar_saldo(dados, conta_dest_reemb, item["valor"], "somar")
                            
                            salvar_dados_usuario(usuario_atual, dados)
                            st.success("Reembolso confirmado!")
                            st.rerun()
                else:
                    # Se já está pago, botão para reabrir/desmarcar
                    if col_e_st.button("Reabrir", key=f"reabrir_{idx}"):
                        # Deduz do saldo onde havia entrado
                        conta_dest_reemb = item.get("conta_destino", "Nu")
                        alterar_saldo(dados, conta_dest_reemb, item["valor"], "subtrair")
                        
                        dados["gastos_terceiros_emprestimo"][idx]["pago"] = False
                        dados["gastos_terceiros_emprestimo"][idx]["forma_recebimento"] = ""
                        dados["gastos_terceiros_emprestimo"][idx]["conta_destino"] = ""
                        
                        salvar_dados_usuario(usuario_atual, dados)
                        st.rerun()

# ----------------- ABA 6: NOVA ABA SALDOS (CONTAS DINÂMICAS) -----------------
with abas[5]:
    st.markdown("<h3 class='titulo-secao'>💳 Gestão de Contas & Saldos</h3>", unsafe_allow_html=True)
    st.caption("Crie, edite, exclua e faça transferências entre suas contas normais ou de poupança (Guardado).")
    
    # 1. Visualizar Saldos em formato de Cards
    st.markdown("#### Seus Saldos Atuais")
    
    col_c1, col_c2 = st.columns(2)
    with col_c1:
        st.markdown("**Contas Correntes (Dia a Dia):**")
        for conta in dados["contas"]:
            if conta.get("tipo", "Normal") == "Normal":
                st.info(f"🏦 **{conta['nome']}:** R$ {conta['saldo']:,.2f}")
    with col_c2:
        st.markdown("**Poupança & Guardado (Investimentos):**")
        for conta in dados["contas"]:
            if conta.get("tipo", "Normal") == "Guardado":
                st.success(f"🔒 **{conta['nome']}:** R$ {conta['saldo']:,.2f}")
                
    st.markdown("---")
    
    # Seções expansíveis para gerenciamento
    aba_ger_conta = st.tabs(["💸 Transferir", "➕ Criar Conta", "⚙️ Renomear/Excluir"])
    
    # 6.1 Lançamentos entre Contas (Transferências)
    with aba_ger_conta[0]:
        st.markdown("**Fazer Transferência/Lançamento entre Contas**")
        with st.form("form_transferencia", clear_on_submit=True):
            contas_nomes = [c["nome"] for c in dados["contas"]]
            origem = st.selectbox("Conta de Origem (Sai o dinheiro):", contas_nomes)
            destino = st.selectbox("Conta de Destino (Entra o dinheiro):", contas_nomes)
            valor_transf = st.number_input("Valor da Transferência (R$)", min_value=0.0, step=50.0, format="%.2f")
            
            confirmar = st.form_submit_button("Realizar Transferência 🔁")
            if confirmar:
                if origem == destino:
                    st.error("A conta de origem não pode ser igual à conta de destino!")
                elif valor_transf <= 0:
                    st.error("O valor deve ser maior do que zero!")
                else:
                    # Processa a transferência
                    if alterar_saldo(dados, origem, valor_transf, "subtrair"):
                        alterar_saldo(dados, destino, valor_transf, "somar")
                        
                        # Salva registro no histórico de transferências
                        nova_transf = {
                            "data": str(datetime.now().date()),
                            "origem": origem,
                            "destino": destino,
                            "valor": valor_transf
                        }
                        dados.setdefault("transferencias", []).append(nova_transf)
                        salvar_dados_usuario(usuario_atual, dados)
                        st.success(f"Transferência de R$ {valor_transf:,.2f} realizada com sucesso!")
                        st.rerun()
                    else:
                        st.error("Erro ao realizar transferência.")

    # 6.2 Criar Nova Conta
    with aba_ger_conta[1]:
        st.markdown("**Adicionar uma Nova Conta ao seu Aplicativo**")
        with st.form("form_criar_conta", clear_on_submit=True):
            nome_nova_conta = st.text_input("Nome da Conta (Ex: Inter, Nubank, Dinheiro Físico):")
            tipo_nova_conta = st.selectbox("Tipo da Conta:", ["Normal", "Guardado"])
            saldo_inicial = st.number_input("Saldo Inicial (R$):", min_value=0.0, step=10.0, format="%.2f")
            
            criar = st.form_submit_button("Criar Nova Conta 🏦")
            if criar and nome_nova_conta:
                # Verificar se o nome já existe
                ja_existe = any(c["nome"].lower() == nome_nova_conta.lower() for c in dados["contas"])
                if ja_existe:
                    st.error("Já existe uma conta cadastrada com esse nome!")
                else:
                    nova_c = {
                        "nome": nome_nova_conta.strip(),
                        "saldo": saldo_inicial,
                        "tipo": tipo_nova_conta
                    }
                    dados["contas"].append(nova_c)
                    salvar_dados_usuario(usuario_atual, dados)
                    st.success(f"Conta '{nome_nova_conta}' criada com sucesso!")
                    st.rerun()

    # 6.3 Renomear ou Excluir Conta
    with aba_ger_conta[2]:
        st.markdown("**Renomear ou Excluir uma Conta Existente**")
        
        contas_nomes = [c["nome"] for c in dados["contas"]]
        conta_selecionada = st.selectbox("Selecione a Conta para alterar:", contas_nomes)
        
        col_ren, col_exc = st.columns(2)
        
        with col_ren:
            st.markdown("**Renomear Conta:**")
            novo_nome_conta = st.text_input("Novo Nome da Conta:", value=conta_selecionada)
            if st.button("Salvar Novo Nome"):
                if novo_nome_conta and novo_nome_conta != conta_selecionada:
                    # Atualizar o nome da conta
                    for c in dados["contas"]:
                        if c["nome"] == conta_selecionada:
                            c["nome"] = novo_nome_conta.strip()
                    
                    # Atualizar referências no histórico de receitas, gastos e transferências
                    for r in dados.get("receitas", []):
                        if r.get("conta") == conta_selecionada:
                            r["conta"] = novo_nome_conta
                    for g in dados.get("gastos_fixos", []):
                        if g.get("conta") == conta_selecionada:
                            g["conta"] = novo_nome_conta
                    for a in dados.get("gastos_avulsos", []):
                        if a.get("conta") == conta_selecionada:
                            a["conta"] = novo_nome_conta
                    for e in dados.get("gastos_terceiros_emprestimo", []):
                        if e.get("conta_origem") == conta_selecionada:
                            e["conta_origem"] = novo_nome_conta
                        if e.get("conta_destino") == conta_selecionada:
                            e["conta_destino"] = novo_nome_conta
                    for t in dados.get("transferencias", []):
                        if t.get("origem") == conta_selecionada:
                            t["origem"] = novo_nome_conta
                        if t.get("destino") == conta_selecionada:
                            t["destino"] = novo_nome_conta
                            
                    salvar_dados_usuario(usuario_atual, dados)
                    st.success("Conta renomeada com sucesso!")
                    st.rerun()
                    
        with col_exc:
            st.markdown("**Excluir Conta:**")
            st.warning("⚠️ Atenção: Ao excluir, os saldos não serão transferidos.")
            if st.button("Confirmar Exclusão ❌"):
                if len(dados["contas"]) <= 1:
                    st.error("Você deve manter pelo menos uma conta ativa no aplicativo!")
                else:
                    dados["contas"] = [c for c in dados["contas"] if c["nome"] != conta_selecionada]
                    salvar_dados_usuario(usuario_atual, dados)
                    st.success(f"Conta '{conta_selecionada}' excluída!")
                    st.rerun()

# ----------------- ABA 7: CONTA ISOLADA & CAIXINHAS (POUPANÇA/GUARDADO) -----------------
with abas[6]:
    st.markdown("<h3 class='titulo-secao'>🔒 Dinheiro Guardado & Poupança</h3>", unsafe_allow_html=True)
    st.markdown("""
    <div class="card-didatico">
        💡 <b>Como funciona seu Dinheiro Guardado:</b><br>
        Toda conta que você cria na aba anterior e define como do tipo <b>"Guardado"</b> (como o Banco Pan ou as Caixinhas) 
        é listada aqui de forma integrada! Para guardar dinheiro, basta ir na aba anterior e fazer uma transferência de uma conta corrente normal para a sua conta de poupança.
    </div>
    """, unsafe_allow_html=True)
    
    # Calcular e mostrar saldos de investimentos
    contas_guardado = [c for c in dados["contas"] if c.get("tipo", "Normal") == "Guardado"]
    total_reservas = sum(c["saldo"] for c in contas_guardado)
    
    st.markdown("#### Detalhamento das Economias:")
    for conta in contas_guardado:
        st.info(f"💰 **{conta['nome']}:** R$ {conta['saldo']:,.2f}")
        
    st.markdown(f"### 📈 Total Acumulado Guardado: **R$ {total_reservas:,.2f}**")
    
    # Histórico de transferências para poupança (economias)
    transferencias_poupanca = []
    for t in dados.get("transferencias", []):
        dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
        orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
        if dest_tipo == "Guardado" and orig_tipo == "Normal":
            transferencias_poupanca.append(t)
            
    if transferencias_poupanca:
        st.markdown("#### Histórico Recente de Economias Guardadas")
        df_dep = pd.DataFrame(transferencias_poupanca)
        df_dep = df_dep[["data", "origem", "destino", "valor"]]
        df_dep.columns = ["Data", "Origem", "Destino (Guardado)", "Valor Guardado (R$)"]
        st.dataframe(df_dep, use_container_width=True)

# ----------------- ABA 8: ALTERAÇÃO DE SENHA (PIN) -----------------
with abas[7]:
    st.markdown("<h3 class='titulo-secao'>⚙️ Alterar Sua Senha (PIN)</h3>", unsafe_allow_html=True)
    st.caption("Sua senha garante a privacidade dos seus dados financeiros se outras pessoas usarem o mesmo aplicativo.")
    
    with st.form("form_alterar_senha"):
        senha_atual = st.text_input("Digite o PIN atual (6 dígitos):", type="password", max_chars=6)
        nova_senha = st.text_input("Digite o NOVO PIN (6 dígitos numéricos):", type="password", max_chars=6)
        nova_senha_conf = st.text_input("Confirme o NOVO PIN (6 dígitos numéricos):", type="password", max_chars=6)
        
        atualizar_btn = st.form_submit_button("Alterar Senha")
        if atualizar_btn:
            if senha_atual != dados.get("pin", "123456"):
                st.error("PIN atual incorreto!")
            elif not nova_senha.isdigit() or len(nova_senha) != 6:
                st.error("O novo PIN deve conter exatamente 6 números.")
            elif nova_senha != nova_senha_conf:
                st.error("A confirmação da nova senha está diferente!")
            else:
                dados["pin"] = nova_senha
                salvar_dados_usuario(usuario_atual, dados)
                st.success("Sua senha (PIN) de 6 dígitos foi alterada com sucesso!")
                st.rerun()
